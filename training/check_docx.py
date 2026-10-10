"""Structural check of an exported question-paper DOCX + preview comparison.

EMU is python-docx's native unit; 1 twip = 635 EMU.
"""
import json
import os
import re
import sys
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.oxml.ns import qn

# ElementTree otherwise serializes the WordprocessingML namespace as the
# auto-generated "ns0" prefix, which breaks the border/table classification
# below. Pin the canonical "w" prefix so the XML matches what Word writes.
ET.register_namespace('w', 'http://schemas.openxmlformats.org/wordprocessingml/2006/main')

A4_W = 11906
A4_H = 16838
MARGIN_SIDE = 720
MARGIN_TOP = 1080
CONTENT_W = A4_W - MARGIN_SIDE * 2
EMU_PER_TWIP = 635


def twips(length):
    return int(length) // EMU_PER_TWIP


def docx_blocks(doc):
    """Ordered block signature mirroring what the A4 preview renders."""
    blocks = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag
        if tag == qn('w:p'):
            # Reconstruct paragraph text in document order, mapping <w:tab/>
            # to "\t" so a section heading and its tab-separated marks equation
            # split exactly like the browser preview does.
            pieces = []
            for run in child.findall('.//' + qn('w:r')):
                for el in run:
                    if el.tag == qn('w:t'):
                        pieces.append(el.text or '')
                    elif el.tag == qn('w:tab'):
                        pieces.append('\t')
            text = ''.join(pieces)
            has_drawing = child.find('.//' + qn('w:drawing')) is not None
            if has_drawing and not text.strip():
                blocks.append({'type': 'picture', 'text': ''})
                continue
            if has_drawing:
                blocks.append({'type': 'picture', 'text': text.strip()})
                continue
            if not text.strip():
                blocks.append({'type': 'blank'})
                continue
            pPr = child.find(qn('w:pPr'))
            keep = pPr is not None and pPr.find(qn('w:keepNext')) is not None
            jc = pPr.find(qn('w:jc')) if pPr is not None else None
            centred = jc is not None and jc.get(qn('w:val')) == 'center'
            if keep and centred:
                blocks.append({'type': 'title', 'text': text.strip()})
            elif keep:
                parts = text.split('\t')
                blocks.append({
                    'type': 'section',
                    'text': parts[0].strip(),
                    'marks': parts[1].strip() if len(parts) > 1 else ''
                })
            else:
                blocks.append({'type': 'text', 'text': text.strip()})
        elif tag == qn('w:tbl'):
            xml = ET.tostring(child, encoding='unicode')
            # Borderless table cells that hold images are the side-by-side
            # diagram grid; surface each drawing as one picture block so it
            # matches the on-screen picture count.
            if child.find('.//' + qn('w:drawing')) is not None and 'w:val="single"' not in xml:
                for p in child.findall('.//' + qn('w:p')):
                    if p.find('.//' + qn('w:drawing')) is not None:
                        blocks.append({'type': 'picture', 'text': ''})
                continue
            rows = []
            for tr in child.findall(qn('w:tr')):
                cells = []
                for tc in tr.findall(qn('w:tc')):
                    cells.append(''.join(t.text or '' for t in tc.findall('.//' + qn('w:t'))).strip())
                rows.append(cells)
            kind = 'table' if 'w:val="single"' in xml else 'options'
            blocks.append({'type': kind, 'rows': rows})
    return blocks


def check(path, preview_path=None):
    report = {"file": path}
    problems = []

    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        report["parts"] = len(names)
        report["media"] = sorted(n for n in names if n.startswith("word/media/") and not n.endswith("/"))
        for part in [n for n in names if n.endswith(".xml") or n.endswith(".rels")]:
            try:
                ET.fromstring(z.read(part))
            except ET.ParseError as e:
                problems.append(f"invalid XML in {part}: {e}")

        document_xml = z.read("word/document.xml").decode("utf-8")
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")
        used_rids = set(re.findall(r'r:embed="(rIdPic\d+)"', document_xml))
        defined_rids = set(re.findall(r'Id="(rIdPic\d+)"', rels))
        if used_rids - defined_rids:
            problems.append(f"embedded rids without relationship: {sorted(used_rids - defined_rids)}")
        for target in re.findall(r'Target="media/([^"]+)"', rels):
            if f"word/media/{target}" not in names:
                problems.append(f"missing media file {target}")
        report["dangling"] = sorted(used_rids - defined_rids)
        report["right_tab_pos"] = 'w:pos="10466"' in document_xml
        content_types = z.read("[Content_Types].xml").decode("utf-8")
        report["has_footer_part"] = "word/footer1.xml" in names
        report["has_footer_rel"] = 'Target="footer1.xml"' in rels
        report["has_footer_override"] = "/word/footer1.xml" in content_types
        report["footer_reference"] = "w:footerReference" in document_xml
        footer_xml = z.read("word/footer1.xml").decode("utf-8") if "word/footer1.xml" in names else ""
        report["footer_page_field"] = 'w:instr=" PAGE "' in footer_xml
        report["has_header_part"] = "word/header1.xml" in names
        report["has_header_rel"] = 'Target="header1.xml"' in rels
        report["has_header_override"] = "/word/header1.xml" in content_types
        report["header_reference"] = "w:headerReference" in document_xml
        header_xml = z.read("word/header1.xml").decode("utf-8") if "word/header1.xml" in names else ""
        header_embed = re.search(r'r:embed="(rIdHeaderImg)"', header_xml) is not None
        report["header_banner_image"] = 'w:drawing' in header_xml and 'rIdHeaderImg' in header_xml
        report["page_breaks"] = document_xml.count("<w:pageBreakBefore/>")
        if not report["has_footer_part"]:
            problems.append("missing footer1.xml part")
        if not report["has_footer_rel"]:
            problems.append("footer1.xml has no relationship")
        if not report["has_footer_override"]:
            problems.append("footer1.xml missing content-type override")
        if not report["footer_reference"]:
            problems.append("sectPr has no footerReference")
        if footer_xml and not report["footer_page_field"]:
            problems.append("footer has no PAGE field")
        if not report["has_header_part"]:
            problems.append("missing header1.xml part")
        if not report["has_header_rel"]:
            problems.append("header1.xml has no relationship")
        if not report["has_header_override"]:
            problems.append("header1.xml missing content-type override")
        if not report["header_reference"]:
            problems.append("sectPr has no headerReference")
        if report["has_header_part"]:
            header_rels_name = "word/_rels/header1.xml.rels"
            if header_rels_name not in names:
                problems.append("header1.xml has no rels part")
            else:
                hr = z.read(header_rels_name).decode("utf-8")
                target = re.search(r'Target="media/([^"]+)"', hr)
                if 'Id="rIdHeaderImg"' not in hr or not target:
                    problems.append("header has no banner image relationship")
                elif f"word/media/{target.group(1)}" not in names:
                    problems.append(f"missing header banner media {target.group(1)}")
                if 'r:embed="rIdHeaderImg"' not in header_xml:
                    problems.append("header does not draw the banner image")

    doc = Document(path)
    sec = doc.sections[0]
    report["page_twips"] = [twips(sec.page_width), twips(sec.page_height)]
    report["margins_twips"] = {side: twips(getattr(sec, f"{side}_margin"))
                               for side in ("top", "bottom", "left", "right")}
    if report["page_twips"] != [A4_W, A4_H]:
        problems.append(f"page size {report['page_twips']} != A4 portrait {A4_W}x{A4_H}")
    for side, val in report["margins_twips"].items():
        limit = MARGIN_TOP if side == "top" else MARGIN_SIDE
        if val > limit:
            problems.append(f"{side} margin {val} > {limit}")
        if val == 0:
            problems.append(f"{side} margin is 0")

    report["paragraphs"] = len(doc.paragraphs)
    report["tables"] = len(doc.tables)
    report["inline_shapes"] = len(doc.inline_shapes)

    for i, table in enumerate(doc.tables):
        width = sum(twips(c.width) for c in table.rows[0].cells)
        if width > CONTENT_W + 20:
            problems.append(f"table {i} width {width} > content width {CONTENT_W}")

    blocks = docx_blocks(doc)
    report["blocks"] = blocks
    report["block_counts"] = {}
    for b in blocks:
        report["block_counts"][b["type"]] = report["block_counts"].get(b["type"], 0) + 1

    # Document must not start with blank pages of empty paragraphs.
    lead = 0
    while lead < len(blocks) and blocks[lead]["type"] == "blank":
        lead += 1
    if lead > 2:
        problems.append(f"{lead} leading blank paragraphs")

    if preview_path and os.path.exists(preview_path):
        with open(preview_path, encoding="utf-8") as fh:
            preview = json.load(fh)
        report["preview_compare"] = compare(preview.get("blocks", []), blocks)
        problems += [f"preview/docx mismatch: {p}" for p in report["preview_compare"]["problems"]]

    report["problems"] = problems
    report["ok"] = not problems
    return report


def _norm(value):
    """Collapse runs of whitespace so the browser-normalized preview text and
    the raw Word paragraph text compare equal."""
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _sig(kind, b):
    if kind == "section":
        return (_norm(b["text"]), _norm(b.get("marks", "")))
    if kind in ("text", "title", "picture"):
        return _norm(b["text"])
    if kind in ("table", "options"):
        return tuple(tuple(_norm(c) for c in r) for r in b["rows"])
    return ()


def compare(preview_blocks, docx_blocks_):
    """Compare the on-screen A4 blocks with the exported Word blocks."""
    problems = []
    tally = lambda blocks, key: {t: sum(1 for b in blocks if b["type"] == t)
                                 for t in set(b["type"] for b in blocks)}
    p_counts, d_counts = tally(preview_blocks, "t"), tally(docx_blocks_, "t")

    for kind in ("title", "section", "picture", "table", "options"):
        if p_counts.get(kind, 0) != d_counts.get(kind, 0):
            problems.append(f"{kind} count preview={p_counts.get(kind, 0)} docx={d_counts.get(kind, 0)}")

    def ordered(blocks, kind, with_marks=False):
        return [_sig(kind, b) for b in blocks if b["type"] == kind]

    for kind in ("title", "section"):
        p, d = ordered(preview_blocks, kind), ordered(docx_blocks_, kind)
        if p != d:
            problems.append(f"{kind} text differs preview={p[:4]} docx={d[:4]}")

    for kind in ("table", "options"):
        p = ordered(preview_blocks, kind)
        d = ordered(docx_blocks_, kind)
        p_shapes = [tuple(len(r) for r in rows) for rows in p]
        d_shapes = [tuple(len(r) for r in rows) for rows in d]
        if p_shapes != d_shapes:
            problems.append(f"{kind} shapes preview={p_shapes} docx={d_shapes}")

    p_text = [_norm(b["text"]) for b in preview_blocks if b["type"] == "text" and b["text"]]
    d_text = [_norm(b["text"]) for b in docx_blocks_ if b["type"] == "text" and b["text"]]
    if len(p_text) != len(d_text):
        problems.append(f"text paragraph count preview={len(p_text)} docx={len(d_text)}")
    else:
        diff = [(i, a, b) for i, (a, b) in enumerate(zip(p_text, d_text)) if a != b]
        if diff:
            i, a, b = diff[0]
            problems.append(f"text block {i} differs preview={a[:60]!r} docx={b[:60]!r}")

    return {
        "preview_counts": p_counts,
        "docx_counts": d_counts,
        "text_paragraphs_preview": len(p_text),
        "text_paragraphs_docx": len(d_text),
        "problems": problems
    }


if __name__ == "__main__":
    out = []
    for path in sys.argv[1:]:
        preview = os.path.splitext(path)[0] + ".preview.json"
        out.append(check(path, preview))
    print(json.dumps(out, ensure_ascii=False, indent=2))
    sys.exit(0 if all(r["ok"] for r in out) else 1)
