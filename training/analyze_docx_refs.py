"""Analyze the Word reference papers in training/ for layout geometry.

Reports, per file: page size/orientation, margins, default font/size, the
font-size x bold histogram, alignment histogram, inline-image sizes, table
counts and numbering, plus representative paragraph samples. Output goes to
training/docx_refs.json and a compact summary is printed.

Run:  python training/analyze_docx_refs.py
"""
import glob
import json
import os
import re

from docx import Document
from docx.shared import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
EMU_PER_TWIP = 635
EMU_PER_PT = 12700


def pt(emu):
    return round(int(emu) / EMU_PER_PT, 1) if emu is not None else None


def twips(emu):
    return int(int(emu) // EMU_PER_TWIP) if emu is not None else None


def para_info(p, doc):
    runs = []
    for r in p.runs:
        sz = r.font.size
        if sz is None:
            sz = doc.styles['Normal'].font.size
        runs.append({
            "t": r.text[:90],
            "sz_pt": round(sz.pt, 1) if sz is not None else None,
            "b": bool(r.bold),
        })
    style = p.style.name if p.style is not None else None
    return {"text": p.text.strip()[:120], "style": style,
            "align": str(p.alignment), "runs": runs}


def analyze(path):
    doc = Document(path)
    p = doc.sections[0]
    margins = {
        "top_in": round((pt(p.top_margin) or 0) / 72, 2),
        "bottom_in": round((pt(p.bottom_margin) or 0) / 72, 2),
        "left_in": round((pt(p.left_margin) or 0) / 72, 2),
        "right_in": round((pt(p.right_margin) or 0) / 72, 2),
    }
    sections_entry = {
        "page_in": [round(p.page_width.inches, 2), round(p.page_height.inches, 2)],
        "orientation": "portrait" if p.page_height >= p.page_width else "landscape",
        "margins_in": margins,
    }

    sz_bold = {}
    align = {}
    samples = []
    style_hist = {}
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        st = p.style.name if p.style is not None else "None"
        style_hist[st] = style_hist.get(st, 0) + 1
        align[(p.alignment or "None")] = align.get(str(p.alignment), 0) + 1
        for r in p.runs:
            if not r.text.strip():
                continue
            sz = r.font.size
            if sz is None:
                sz = doc.styles['Normal'].font.size
            key = (round(sz.pt, 1) if sz else None, bool(r.bold))
            sz_bold[key] = sz_bold.get(key, 0) + 1
        if len(samples) < 25:
            samples.append({"style": st, "align": str(p.alignment), "t": text[:100]})

    imgs = []
    for sh in doc.inline_shapes:
        imgs.append({"w_pt": pt(sh.width), "h_pt": pt(sh.height)})

    tables = []
    for t in doc.tables:
        rows = len(t.rows)
        cols = len(t.columns) if t.columns else 0
        tables.append({"rows": rows, "cols": cols, "style": t.style.name if t.style else None})

    numbering = 0
    for p in doc.paragraphs:
        if p._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr') is not None:
            pPr = p._p.pPr
            if pPr is not None and pPr.numPr is not None:
                numbering += 1

    return {
        "sections": sections_entry,
        "default_font": doc.styles['Normal'].font.name,
        "default_size_pt": round(doc.styles['Normal'].font.size.pt, 1) if doc.styles['Normal'].font.size else None,
        "font_bold_hist": {f"{k[0]}pt/{'bold' if k[1] else 'normal'}": v for k, v in sorted(
            sz_bold.items(), key=lambda kv: (kv[0][0] or 0, kv[0][1]))},
        "align_hist": align,
        "styles": style_hist,
        "num_numbered_paras": numbering,
        "images": imgs[:40],
        "image_count": len(imgs),
        "tables": tables,
        "samples": samples,
    }


def main():
    report = {}
    files = sorted(glob.glob(os.path.join(HERE, "*.docx")))
    for path in files:
        fn = os.path.basename(path)
        try:
            info = analyze(path)
            report[fn] = info
            s = info["sections"]
            print(f"{fn}: {s['page_in']} {s['orientation']} margins={s['margins_in']} "
                  f"font={info['default_font']} {info['default_size_pt']}pt "
                  f"imgs={info['image_count']} tables={len(info['tables'])}")
            print(f"   sizes={info['font_bold_hist']}")
            print(f"   styles={info['styles']} align={info['align_hist']}")
        except Exception as e:
            report[fn] = {"error": str(e)}
            print(f"{fn}: ERROR {e}")
    with open(os.path.join(HERE, "docx_refs.json"), "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=1)
    print("wrote docx_refs.json")


if __name__ == "__main__":
    main()