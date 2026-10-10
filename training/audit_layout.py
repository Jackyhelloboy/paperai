"""Measure the current reference corpus without treating unmatched files as pairs.

Run with Python containing python-docx and PyMuPDF:
    python training/audit_layout.py
Only training/*.pdf and training/*.docx are inspected; generated out/ is excluded.
"""
import hashlib
import json
from pathlib import Path
from collections import Counter
import fitz
from docx import Document

ROOT = Path(__file__).resolve().parent

def audit():
    records = []
    seen = {}
    for path in sorted(ROOT.iterdir()):
        if path.suffix.lower() not in {'.pdf', '.docx'}:
            continue
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        record = {'file': path.name, 'sha256': sha, 'bytes': path.stat().st_size}
        if sha in seen:
            record['duplicate_of'] = seen[sha]
        seen.setdefault(sha, path.name)
        if path.suffix.lower() == '.pdf':
            with fitz.open(path) as doc:
                counts = [len(p.get_text().strip()) for p in doc]
                record.update(kind='scan-input' if not any(counts) else 'printed-reference',
                    pages=len(doc), text_chars=sum(counts),
                    page_points=[[round(p.rect.width,2), round(p.rect.height,2)] for p in doc],
                    image_counts=[len(p.get_images()) for p in doc])
        else:
            doc = Document(path)
            fonts, sizes = Counter(), Counter()
            for paragraph in doc.paragraphs:
                for run in paragraph.runs:
                    if not run.text.strip(): continue
                    if run.font.name: fonts[run.font.name] += len(run.text)
                    if run.font.size: sizes[str(run.font.size.pt)] += len(run.text)
            excluded = path.name.startswith('final PT_')
            record.update(kind='excluded-non-question-paper' if excluded else 'word-reference',
                paragraphs=len(doc.paragraphs), tables=len(doc.tables), inline_images=len(doc.inline_shapes),
                sections=[{'page_inches':[round(s.page_width.inches,3),round(s.page_height.inches,3)],
                    'margins_inches':{k:round(getattr(s,k+'_margin').inches,3) for k in ['top','bottom','left','right']}}
                    for s in doc.sections], explicit_fonts=dict(fonts), explicit_sizes=dict(sizes))
        records.append(record)
    return {'version':4, 'scope':'Top-level source documents only; no supervised input/output pairs established.',
        'output_policy':'A4 portrait, 0.5 inch margins; native editable questions/tables/grids/basic shapes; original illustrations cropped separately.',
        'runtime_profile':'frontend/vendor/paper-layout.js', 'ocr_instructions':'worker/src/paper-layout-rules.js',
        'records':records}

if __name__ == '__main__':
    target = ROOT / 'layout-audit.json'
    target.write_text(json.dumps(audit(),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f'Wrote {target.name}')
